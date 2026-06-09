"""Ephemeral resume parsing adapters.

All parsers operate on bytes passed by the caller and return in-memory text.
No source document bytes are written to disk by this module.
"""

from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from .models import ResumeDocument, SourceBullet

_SECTION_NAMES = {
    "experience",
    "work experience",
    "projects",
    "education",
    "skills",
    "summary",
    "leadership",
}

_BULLET_PREFIX = re.compile(r"^\s*(?:[-*•‣▪]|\d+[.)])\s+")


def parse_resume_document(payload: bytes, filename: str, session_id: str) -> ResumeDocument:
    """Parse a resume payload into a worker-local `ResumeDocument`.

    Supported built-ins are text and Markdown. PDF/DOCX/unstructured support is
    attempted through optional runtime imports so the worker remains testable
    before dependency installation.
    """

    if not payload:
        raise ValueError("resume payload is empty")

    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = payload.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        text = _parse_pdf(payload)
    elif suffix == ".docx":
        text = _parse_docx(payload)
    else:
        text = _parse_unstructured(payload, filename)

    normalized = _normalize_text(text)
    digest = hashlib.sha256(payload).hexdigest()
    document_id = f"resume_{digest[:16]}"
    bullets = tuple(_extract_source_bullets(normalized))
    if not bullets:
        raise ValueError("resume contains no source bullets to ground rewrites")
    return ResumeDocument(
        document_id=document_id,
        session_id=session_id,
        source_name=filename,
        text_sha256=digest,
        raw_text=normalized,
        source_bullets=bullets,
    )


def _parse_pdf(payload: bytes) -> str:
    try:
        import fitz  # type: ignore[import-not-found]

        with fitz.open(stream=payload, filetype="pdf") as document:
            return "\n".join(page.get_text("text") for page in document)
    except Exception as pymupdf_error:
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]

            reader = PdfReader(io.BytesIO(payload))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as pypdf_error:
            raise ValueError(
                "unable to parse PDF with PyMuPDF or pypdf"
            ) from pypdf_error or pymupdf_error


def _parse_docx(payload: bytes) -> str:
    try:
        from docx import Document  # type: ignore[import-not-found]

        document = Document(io.BytesIO(payload))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        table_cells: list[str] = []
        for table in document.tables:
            for row in table.rows:
                table_cells.extend(cell.text for cell in row.cells if cell.text)
        return "\n".join(paragraphs + table_cells)
    except Exception as error:
        raise ValueError("unable to parse DOCX with python-docx") from error


def _parse_unstructured(payload: bytes, filename: str) -> str:
    try:
        from unstructured.partition.auto import partition  # type: ignore[import-not-found]

        elements = partition(file=io.BytesIO(payload), metadata_filename=filename)
        return "\n".join(str(element) for element in elements)
    except Exception as error:
        raise ValueError(f"unsupported resume file type for {filename}") from error


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_source_bullets(text: str) -> list[SourceBullet]:
    current_section = "unknown"
    bullets: list[SourceBullet] = []
    continuation: list[str] = []
    continuation_line = 0

    def flush_continuation() -> None:
        nonlocal continuation, continuation_line
        if continuation:
            joined = " ".join(continuation).strip()
            if _looks_like_resume_claim(joined):
                bullets.append(
                    SourceBullet(
                        id=f"src_{len(bullets) + 1:04d}",
                        section=current_section,
                        text=joined,
                        line_number=continuation_line,
                    )
                )
            continuation = []
            continuation_line = 0

    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            flush_continuation()
            continue
        lowered = line.rstrip(":").lower()
        if lowered in _SECTION_NAMES and len(line.split()) <= 3:
            flush_continuation()
            current_section = lowered
            continue
        if _BULLET_PREFIX.match(line):
            flush_continuation()
            bullet_text = _BULLET_PREFIX.sub("", line).strip()
            bullets.append(
                SourceBullet(
                    id=f"src_{len(bullets) + 1:04d}",
                    section=current_section,
                    text=bullet_text,
                    line_number=line_number,
                )
            )
            continue
        if _looks_like_resume_claim(line):
            if not continuation:
                continuation_line = line_number
            continuation.append(line)
    flush_continuation()
    return bullets


def _looks_like_resume_claim(line: str) -> bool:
    words = line.split()
    if len(words) < 5:
        return False
    verbs = {
        "built",
        "created",
        "designed",
        "delivered",
        "implemented",
        "improved",
        "increased",
        "led",
        "managed",
        "migrated",
        "reduced",
        "shipped",
        "owned",
        "automated",
    }
    lowered = {word.strip(".,:;()").lower() for word in words}
    return bool(lowered & verbs) or any(char.isdigit() for char in line)
